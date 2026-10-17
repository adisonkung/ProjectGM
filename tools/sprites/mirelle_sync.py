"""Generate Mirelle Final Sync JSON/PNG/XLSX and 60FPS asset verification videos.

Input: tools/dev/mirelle-sync.ts export (actual runtime clocks/FX sampling/engine casts).
python3 tools/sprites/mirelle_sync.py dist/mirelle-sync-runtime.json [sheet dir] [video dir]
Original source workbooks/assets remain read-only. This renderer is an asset sync preview, not a game screenshot.
"""
import hashlib
import json
import math
import os
from pathlib import Path
import subprocess
import sys
import tempfile

from PIL import Image, ImageDraw, ImageFont
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
D = json.loads(Path(sys.argv[1]).read_text())
OUT = Path(sys.argv[2] if len(sys.argv) > 2 else 'docs/timelines/mirelle')
VIDEOS = Path(sys.argv[3] if len(sys.argv) > 3 else 'tmp/mirelle-final-sync')
OUT.mkdir(parents=True, exist_ok=True)
VIDEOS.mkdir(parents=True, exist_ok=True)
FONT = '/usr/share/fonts/opentype/tlwg/Loma.otf'
BOLD = '/usr/share/fonts/opentype/tlwg/Loma-Bold.otf'
if not Path(FONT).exists():
    # The installed Noto Thai subset omits Latin. Merge installed families for bilingual sheet labels.
    from fontTools.merge import Merger
    FONT_DIR = tempfile.TemporaryDirectory(prefix='mirelle-sync-fonts-')
    paths = []
    for weight in ['Regular', 'Bold']:
        merged = Merger().merge([f'/usr/share/fonts/truetype/noto/NotoSans-{weight}.ttf',
                                f'/usr/share/fonts/truetype/noto/NotoSansThai-{weight}.ttf'])
        path = str(Path(FONT_DIR.name) / f'Sync-{weight}.ttf')
        merged.save(path)
        paths.append(path)
    FONT, BOLD = paths
font = lambda n: ImageFont.truetype(FONT, n)
FPS = D['fps']
CAN = D['canvas']
W, H, G, X, CELL = CAN['width'], CAN['height'], CAN['ground'], CAN['rootX'], CAN['cell']
CACHE = {}


def png(name):
    if name not in CACHE:
        CACHE[name] = Image.open(ROOT / 'public/sprites' / name).convert('RGBA')
    return CACHE[name]


def tile(name, col, row, w, h):
    return png(name).crop((col*w, row*h, (col+1)*w, (row+1)*h))


def at_cell(at):
    return X + (at['c']-D['caster']['c'])*CELL, G + (at['r']-D['caster']['r'])*CELL


def draw_fx(im, effect):
    cfg = D['fx'][effect['name']]
    part = tile(cfg['file'], effect['frame'], 0, cfg['w'], cfg['h'])
    part = part.resize((max(1, round(effect['width']*2)), max(1, round(effect['height']*2))), Image.Resampling.LANCZOS)
    alpha = effect.get('alpha', 1)
    if alpha < 1:
        part.putalpha(part.getchannel('A').point(lambda a: round(a*alpha)))
    pivot = cfg.get('pivot', [.5, .5])
    im.alpha_composite(part, (round((effect['x']-effect['width']*pivot[0])*2), round((effect['y']-effect['height']*pivot[1])*2)))


def render(skill, destination):
    # Render at device scale2 directly; no CSS-size capture enlarged afterwards.
    process = subprocess.Popen(['ffmpeg', '-v', 'error', '-y', '-f', 'rawvideo', '-pix_fmt', 'rgb24',
        '-s', f'{W*2}x{H*2}', '-r', str(FPS), '-i', '-', '-an', '-c:v', 'libx264', '-crf', '16',
        '-pix_fmt', 'yuv420p', '-movflags', '+faststart', str(destination)], stdin=subprocess.PIPE)
    enemy = D['monster']
    es = enemy['sheet']
    ew, eh = es.get('cellW', 132), es.get('cellH', 100)
    idle = es['anims']['idle']
    for fr in skill['frames']:
        im = Image.new('RGBA', (W*2, H*2), '#101a29')
        dr = ImageDraw.Draw(im)
        for c in range(-1, 9):
            x = round((X-CELL/2+c*CELL)*2)
            dr.line((x, 200, x, H*2-80), fill='#21364c', width=2)
        for r in range(4):
            y = round((G-CELL*.7725+(r-1)*CELL)*2)
            dr.line((0, y, W*2, y), fill='#21364c', width=2)
        for effect in fr['effects']:
            if effect['under']:
                draw_fx(im, effect)
        char = tile(D['sprite'], fr['pose'], fr['row'], D['sheet']['cellW'], D['sheet']['cellH'])
        im.alpha_composite(char.resize((520, 340), Image.Resampling.NEAREST), (round((X-130)*2), round((G-155)*2)))
        for target in (D['targets'][:1] if skill['slot'] == 'skill' else D['targets']):
            x, y = at_cell(target['pos'])
            part = tile(enemy['sprite'], 0, idle['row'], ew, eh)
            im.alpha_composite(part.resize((ew*2, eh*2), Image.Resampling.NEAREST), (round((x-ew/2)*2), round((y-eh+15)*2)))
            statuses = fr['statuses'].get(str(target['id']), [])
            label = ' / '.join(v['id'] for v in statuses)
            dr.text((int(x*2-80), int(y*2+10)), label, font=font(16), fill='#b3e7ff')
        for effect in fr['effects']:
            if not effect['under']:
                draw_fx(im, effect)
        if fr.get('projectile'):
            draw_fx(im, fr['projectile'])
        dr.text((24, 18), f"FINAL SYNC | Mirelle | {skill['label']} | T{fr['T']:03d} | {fr['t']:.3f}s", font=font(24), fill='#e3efff')
        dr.text((24, 58), 'Asset verification playback / real-engine event schedule / training targets', font=font(18), fill='#96aec5')
        live = [event_name(e['event']) for e in skill['events'] if 0 <= fr['t']-e['t'] < .2]
        dr.text((24, H*2-58), ' | '.join(dict.fromkeys(live))[:130], font=font(20), fill='#ffdb86')
        process.stdin.write(im.convert('RGB').tobytes())
    process.stdin.close()
    if process.wait():
        raise RuntimeError('ffmpeg failed')


def event_name(e):
    if e['type'] == 'animEvent':
        return 'ACTIVATION_SCHEDULER' if e['event'] == 'DAMAGE' else e['event']
    if e['type'] == 'damage':
        return 'DAMAGE (contact / field pulse)'
    if e['type'] == 'buff':
        return 'STATUS ' + e['buff']
    if e['type'] == 'impact':
        return 'CONTACT' if e['kind'] == 'skill' else 'EFFECT_RESOLVED (no added burst)'
    return e['type'].upper()


def make_sheet(skill):
    slot = {'skill': 's1', 'skill2': 's2', 'ult': 'ult'}[skill['slot']]
    name = f'mirelle_{slot}_final_sync'
    workbook = next((ROOT/'art/incoming/mirelle'/slot).glob('*.xlsx'))
    sha = hashlib.sha256(workbook.read_bytes()).hexdigest()
    frames = skill['frames']
    n = len(frames)
    count = skill['clip']['frames']
    char = []
    for fr in frames:
        c = fr['char']
        if not char or char[-1]['char'] != f'C{c:02d}':
            char.append({'from': fr['T'], 'to': fr['T'], 'char': f'C{c:02d}'})
        else:
            char[-1]['to'] = fr['T']
    events = []
    for item in skill['events']:
        e, t = item['event'], item['t']
        if e['type'] in ['cast', 'projectile', 'damage', 'buff', 'impact', 'FIELD_STARTED', 'FIELD_ENDED', 'animEvent']:
            event = event_name(e)
            note = (f"Combat event at {t*1000:.3f}ms; " + ('visual ending only' if e['type']=='FIELD_ENDED' else ''))
            row = [round(t*FPS), event, note]
            if row not in events and row[0] < n:
                events.append(row)
    tracks = list(dict.fromkeys(s['name'] for s in skill['scheduled']))
    if slot == 's1':
        tracks.insert(0, 'mf-s1-projectile')
    vfx = []
    for fr in frames:
        for effect in fr['effects'] + ([fr['projectile']] if fr.get('projectile') else []):
            vfx.append({'T': fr['T'], 'track': effect['name'], 'frame': effect['frame']+1,
                'opacity': effect.get('alpha',1), 'split': 'behind' if effect.get('under') else 'front'})
    hitT = round(skill['timing']['hit']*FPS)
    recovery = next(e['t'] for e in skill['nominalEvents'] if e['event']=='RECOVERY_START')
    recoveryT, endT = math.ceil(recovery*FPS-1e-8), math.ceil(skill['timing']['total']*FPS-1e-8)
    phases = [[0, hitT-1, 'A', 'ANTICIPATION', [90,130,200], 'ตั้งท่า / รวมพลัง'],
        [hitT, recoveryT-1, 'B', 'RELEASE', [85,190,230], 'ปล่อยเวท'],
        [recoveryT, endT-1, 'C', 'RECOVERY', [120,185,140], 'คืนท่า'],
        [endT,n-1,'D','INDEPENDENT EFFECTS',[155,145,210],'สนาม / สถานะหลังจบท่า']]
    if slot == 'ult':
        fieldEndT = round(next(e['t'] for e in skill['events'] if e['event']['type']=='FIELD_ENDED')*FPS)
        phases[-1] = [endT,fieldEndT-1,'D','FIELD ACTIVE',[110,180,210],'สนามคงอยู่']
        phases.append([fieldEndT,n-1,'E','SHATTER / RESIDUAL',[200,150,110],'ผลึกแตก / จางหาย'])
    poses = {str(i): [f'F{i}', 'Approved source pose', ''] for i in range(1,count+1)}
    poses[str(count+1)] = ['IDLE', 'Idle cycle after cast; no cast-lock', 'Independent field/projectile/status lifecycle']
    colours = [[120,200,255],[100,220,205],[210,185,255],[255,210,120],[170,185,205]]
    metadata = {'name':name, 'title':f'Final Timeline Sync — Mirelle · {skill["label"]}',
        'subtitle':f'60 FPS sampling · {n} frames · exact runtime ms in Character / Nominal vs Runtime tabs · user 2026-10-17',
        'phases':phases,'poses':poses,'gameSteps':skill['clip']['steps']+[0], 'hits':[round(t*FPS) for t in skill['pulseTimes']],
        'events':events,'thumbs':sorted(set([0,min(12,n-1),hitT,recoveryT,endT,min(n-1,fieldEndT) if slot=='ult' else min(n-1,round(skill['pulseTimes'][0]*FPS)),n-1])),
        'tracks':[[key,key,colours[i%len(colours)],D['fx'][key]['file'], 'L2 behind units' if D['fx'][key].get('layer')=='under' else 'L4 front',
            'fixed field centre' if slot=='ult' and key!='mf-ult-charge' else 'caster / target / flight, runtime placement', 'Runtime fxPoseAt / field lifecycle'] for i,key in enumerate(tracks)],
        'trackNames':{key:key.replace('mf-','') for key in tracks},
        'layers':[[1,'Grid / background'],[2,'Ground/charge/field under units'],[3,'Character + training monsters'],[4,'Projectile / contact / shatter'],[5,'Status and verification labels']],
        'layerNote':'Source art remains separate. Fixed field position; shatter is visual only. No new gameplay layers.',
        'travelLabel':'Root travel (px)', 'heightLabel':'Root height (px)', 'shake':{},
        'notes':['Exact milliseconds: runtime steps/events; chart samples60FPS. Combat applies events in50ms steps.',
            'S1 contact follows actual projectile arrival; launch marker does not apply damage/Slow.',
            'S2 uses pre-hit Chill snapshot. Ult:6 light pulses, no launch/shatter burst; loop survives recovery.',
            'Reviewed sizes/approved art retained. Clips0.9/1.8/1.8s; action timing finalized by user request.'],
        'readme':[f'FINAL Timeline Sync — Mirelle {skill["label"]}',f'Source workbook: {workbook.name}',f'SHA256: {sha}',
            'User authorized finalization 2026-10-17 and slightly longer weighted motion for smoother recovery.',
            'Original workbook remains read-only; S2/Ult10-frame placeholders mapped to actual20 poses.',
            'S1 releaseF7/500ms; S2 activateF13/1000ms; Ult activateF15/1200ms.',
            '60FPS tables sample exact authored ms; the engine resolves on its existing50ms combat step.',
            'Fixtures: distance4cells, first target pre-Chilled, high-HP training monsters. Not final balance.',
            'Video is synchronized asset playback using runtime sampling/engine event traces, not a game screenshot.'],
        'gameNote':'Game ms = exact authored steps; last IDLE row0 is not part of the cast. Preview includes independent effects/statuses.'}
    doc = {'schema':1,'status':'FINAL_SYNC','fps':FPS,'frames':n,'ground':G,
        'sourceWorkbook':{'file':str(workbook.relative_to(ROOT)),'sha256':sha},
        'runtime':{k:skill[k] for k in ['slot','timing','clip','nominalEvents','events','scheduled','flight','pulseTimes','fieldDuration']},
        'characterTiming':char, 'rootTracks':{'travel':[[0,0],[n-1,0]],'height':[[0,0],[n-1,0]]},
        'perFrame':[{'T':fr['T'],'char':f"C{fr['char']:02d}",'drawnFrame':fr['pose']+1,'root':[X,G],
                    'spearTip':[X,G-CELL*.3525],'center':[X,G-39]} for fr in frames],
        'vfx':vfx,'sheet':metadata,'gameFx':{key:D['fx'][key] for key in tracks}}
    path = OUT/f'{name}.json'
    path.write_text(json.dumps(doc,ensure_ascii=False,indent=1)+'\n')
    video = VIDEOS/f'{name}_60fps.mp4'
    render(skill,video)
    env = {**os.environ,'SYNC_SHEET_FONT':FONT,'SYNC_SHEET_FONT_BOLD':BOLD}
    subprocess.run([sys.executable,str(ROOT/'tools/sprites/ult_sync_sheet.py'),str(path),str(video),str(OUT)],env=env,check=True)
    wb = load_workbook(OUT/f'{name}.xlsx')
    ws = wb.create_sheet('Nominal vs Runtime')
    ws.append(['Semantic event','Nominal ms','Nominal T @60FPS','Observed combat ms','Observed T','Authority / note'])
    for ev in skill['nominalEvents']:
        if ev['event']=='IMPACT':
            observed = skill['pulseTimes'][0] if slot=='s1' else skill['timing']['hit']
            note = 'Actual contact, never the nominal launch marker' if slot=='s1' else 'Actual area contact'
        else:
            matching = next((e for e in skill['events'] if e['event']['type']=='animEvent' and e['event']['event']==ev['event']),None)
            observed = matching['t'] if matching else None
            note = 'Activation scheduler marker only' if ev['event']=='DAMAGE' else '50ms engine step'
        ws.append([ev['event'],ev['t']*1000,ev['t']*FPS,None if observed is None else observed*1000,None if observed is None else observed*FPS,note])
    ws = wb.create_sheet('Source Reconciliation')
    for row in [['Item','Original workbook','Final runtime'],['Character count','10 (S2/Ult placeholders)',str(count)],
        ['Release/activation','F7 / F7 hypothetical / F5 hypothetical',f"F{int(skill['clip']['events'].get('PROJECTILE_SPAWN',skill['clip']['events'].get('SKILL_ACTIVATE')))+1}"],
        ['Clip length','12FPS proposed',f"{skill['timing']['total']*1000:.0f}ms, weighted poses"],
        ['Collision vocabulary','Physics/hitbox proposal','Existing targeted flight scheduler; no new collision/miss system'],
        ['S2 field','Optional/unverified','One-shot2x2, no new DoT'],['Ultimate finisher','Natural-expiry visual','Large visual only, no added damage']]: ws.append(row)
    for ws in [wb['Nominal vs Runtime'],wb['Source Reconciliation']]:
        ws.freeze_panes='A2'
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=36
    wb.save(OUT/f'{name}.xlsx')
    print(name, 'complete', flush=True)


for skill in D['skills']:
    make_sheet(skill)
