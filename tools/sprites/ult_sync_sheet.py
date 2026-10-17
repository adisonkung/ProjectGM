"""Timeline Sync Sheet of a skill preview — Bai Long Nine Dragon Sequence (Preview v3) by default; a timeline JSON
with a `sheet` block (phases / poses / events / tracks / layers / notes, e.g. tools/sprites/yuezheng_s1.py) draws any skill.

Reads the preview's timeline JSON (bailong_ult_preview.py) + its MP4 and writes:
  <out>/ult_sync_sheet_v3.png    X-sheet: key-frame thumbnails · phases · events · character holds · root travel / height
                                 · one row per VFX track (frame number, opacity, layer split) · camera shake · hit-stop
  <out>/ult_sync_sheet_v3.xlsx   every frame T000–T146 + character / track / event / layer tables

Usage: python3 tools/sprites/ult_sync_sheet.py <timeline.json> <preview_60fps.mp4> <out dir>
"""
import json
import os
import subprocess
import sys
import tempfile

from PIL import Image, ImageDraw, ImageFont

TL, MP4, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
D = json.load(open(TL))
N, FPS, CELL_PX = D['frames'], D['fps'], 80

# ------------------------------------------------------------------ authored tables (same numbers as the preview script)
PHASES = [  # (from, to, id, name, colour)
    (0, 24, 'A', 'Charge', (120, 98, 200)),
    (25, 44, 'B', 'Thrust · Hit 1', (70, 140, 220)),
    (45, 64, 'C', 'Sweep · Hit 2', (60, 170, 170)),
    (65, 85, 'D', 'Leap', (90, 170, 90)),
    (86, 95, 'E', 'Air gather', (200, 170, 60)),
    (96, 112, 'F', 'Dive', (220, 120, 50)),
    (113, 130, 'G', 'Ground impact · Hit 3', (215, 70, 70)),
    (131, 146, 'H', 'Residual · Recover', (130, 130, 140)),
]
PHASE_TH = {'A': 'ชาร์จ', 'B': 'แทง (Hit 1)', 'C': 'กวาด (Hit 2)', 'D': 'กระโดด', 'E': 'กลางอากาศ', 'F': 'ดิ่ง',
            'G': 'ปักพื้น (Hit 3)', 'H': 'เศษพลัง / กลับท่า'}
POSE = {  # C → (game frame, pose, key point)
    1: ('F1', 'ท่าพร้อม (คล้าย idle) อาวุธต่ำ', 'ต่อจาก idle'),
    2: ('F2', 'รวมพลัง: ย่อตัว ดึงอาวุธไปหลัง', 'วงชาร์จโต'),
    3: ('F3', 'ย่อสุด / ง้างต่ำ ปลายอาวุธเฉียงลงหน้า', 'ANTICIPATION_PEAK'),
    4: ('F4', 'พุ่งต่ำไปหน้า ตัวเอน เท้าไถ', 'STRIKE_START 1 · เริ่มเคลื่อนไปหน้า'),
    5: ('F4', 'ท่าเดียวกับ C04 (ค้างระหว่างพุ่ง)', 'ใช้ภาพซ้ำ'),
    6: ('F5', 'แทงเหยียดสุด', 'HIT 1 · ปลายอาวุธ = จุดเกิด hint'),
    7: ('F6', 'ดึงกลับ บิดสะโพกเตรียมกวาด', 'STRIKE_START 2'),
    8: ('F7', 'กวาดแนวนอนผ่านหน้า', 'HIT 2 · ปลายอาวุธ = จุดเกิด emerge'),
    9: ('F8', 'follow-through อาวุธไปอยู่หลัง', 'ประกายสัมผัสของ hit 2'),
    10: ('F9', 'ย่อเตรียมกระโดด', ''),
    11: ('F10', 'ถีบตัวขึ้น ขาเหยียดดันพื้น', 'STRIKE_START 3 · ascent เริ่มที่พื้น'),
    12: ('F11', 'ลอยขึ้น เข่างอ', ''),
    13: ('F12', 'ลอย หมุนตัว', 'VFX_SPAWN 1'),
    14: ('F13', 'ใกล้จุดสูงสุด อาวุธยกขึ้น', 'coil เริ่มขดรอบตัว'),
    15: ('F14', 'ยกอาวุธเหนือหัว รวมพลัง (ค้าง)', 'FINAL_ANTICIPATION · สูงสุด 64 px'),
    16: ('F15', 'เริ่มดิ่ง ตัวเอียงไปหน้า อาวุธชี้ลงหน้า', 'VFX_SPAWN 2 · dive เริ่ม'),
    17: ('F16', 'ดิ่งเร็ว ตัวยืด', 'dive จบตอนแตะพื้น T112'),
    18: ('F17', 'ปักอาวุธลงพื้นหน้าตัว เข่าลงพื้น', 'HIT 3 · final impact หน้าตัว 1 ช่อง'),
    19: ('F18', 'ค้างท่าปัก แล้วเริ่มดึงอาวุธขึ้น', 'RECOVERY_START'),
    20: ('F19', 'ลุกกลับท่ายืน', 'ต่อกับ idle'),
}
GAME_STEP_MS = [90, 90, 140, 80, 200, 100, 180, 100, 90, 140, 90, 90, 90, 200, 80, 120, 260, 120, 120]
TRACKS = [  # key, label, colour, source, layers, anchor rule, note
    ('charge', 'charge  ground aura', (150, 120, 235), 'fx-bailong-v3-charge.png (10)', 'L2 far half · L6 near half (split lower)',
     'bottom-centre at body x, on the ground', 'f1–f6 grow · f7–f8 skipped (no full dragon yet) · f9–f10 energy shoots into the weapon'),
    ('hint', 'hint  small dragon', (90, 160, 240), 'fx-bailong-v3-hint.png (6)', 'L3 body · L5 head (split tip)',
     'tail (left end) on the weapon tip, follows the tip', 'frame held T033–T035 = hit-stop of hit 1'),
    ('impact', 'tip flash  contact', (240, 240, 255), 'fx-bailong-v3-impact.png f2 ×0.5', 'L6',
     'bottom-centre just under the weapon tip', 'hit 1 contact T031–T036 · hit 2 contact T054–T058'),
    ('emerge', 'emerge  sweep dragon', (70, 200, 190), 'fx-bailong-v3-emerge.png (8)', 'L3 · head 40 % on L5 during T055–T060',
     'tail on the body, slides forward 1.6 px / frame', 'behind first, head crosses in front of the body at the sweep'),
    ('ascent', 'ascent  rising dragon', (110, 200, 110), 'fx-bailong-v3-ascent.png (8)', 'L3 / L5 band (body passes through)',
     'bottom-centre: ground → follows the rising body', ''),
    ('coilA', 'coil  ring around body', (235, 200, 80), 'v3_coilA.png (10, loop)', 'L3 body 60 % · L5 head 40 %',
     'centre on the airborne body', ''),
    ('core', 'core  overhead flare', (255, 230, 140), 'fx-bailong-v3-core.png (12, uses f3–f12)', 'L3 · head 40 % L5',
     'centre slightly above the body', ''),
    ('dive', 'dive  head-first dragon', (240, 140, 60), 'fx-bailong-v3-dive.png (8)', 'L3 body · L5 head-bottom',
     'head on the weapon tip, axis along the shaft', 'gone by T112 — hands over to final impact'),
    ('final', 'final impact', (235, 80, 80), 'fx-bailong-v3-final.png (12)', 'L3 far ring · L6 near ring + debris (split lower)',
     'bottom-centre on G = weapon contact + 1 cell (80 px)', 'T113 flash · T115 core · T117 column · T120 ring · f6–f8 debris · f9–f12 decay'),
    ('residual', 'residual  sparks', (170, 170, 185), 'fx-bailong-v3-residual.png (6, uses f3–f6)', 'L3 (behind)',
     'centre near G, 40 px up', 'opacity × 0.8 and fades to 0 by T146'),
]
TRACK_TH = {'charge': 'ออร่าชาร์จที่พื้น', 'hint': 'มังกรเล็กจากปลายอาวุธ', 'impact': 'ประกายจุดสัมผัส', 'emerge': 'มังกรตามวงกวาด',
            'ascent': 'มังกรหมุนขึ้น', 'coilA': 'มังกรขดรอบตัว', 'core': 'แสงรวมเหนือหัว', 'dive': 'มังกรดิ่ง', 'final': 'ระเบิดพื้นสุดท้าย',
            'residual': 'เศษพลัง'}
EVENTS = [  # T, game event, note
    (0, 'CAST_START', 'เริ่มท่า · charge fade-in 4 เฟรม'),
    (8, 'CHARGE_START', 'C02 ย่อรวมพลัง'),
    (15, 'ANTICIPATION_PEAK', 'C03 ง้างต่ำสุด'),
    (20, 'STRIKE_START 1', 'C04 พุ่งต่ำ · ตัวเริ่มเคลื่อนไปหน้า'),
    (30, 'DAMAGE + IMPACT 1 (25 %)', 'C06 แทงสุด · ในเกม hit-stop 30 ms สั่น 2 px'),
    (33, 'HIT-STOP 1', 'preview: ค้างเฟรม hint T033–T035 + สั่นเบา T033–T036'),
    (36, 'STRIKE_START 2', 'C07 ดึงกลับ'),
    (45, 'DAMAGE + IMPACT 2 (30 %)', 'C08 กวาด · ในเกม hit-stop 40 ms สั่น 4 px · ประกายสัมผัส T054–T058'),
    (65, 'STRIKE_START 3', 'C11 ถีบตัว · ascent เริ่ม'),
    (79, 'VFX_SPAWN 1', 'C13 ลอยหมุน'),
    (86, 'COIL', 'coil เริ่มรอบตัวกลางอากาศ'),
    (91, 'FINAL_ANTICIPATION', 'C15 ยกอาวุธเหนือหัว · จุดสูงสุด 64 px'),
    (94, 'CORE', 'แสงรวมเหนือหัว'),
    (96, 'VFX_SPAWN 2', 'C16 เริ่มดิ่ง'),
    (101, 'DIVE', 'มังกรดิ่ง (จบ T111)'),
    (112, 'GROUND CONTACT', 'ความสูง = 0'),
    (113, 'DAMAGE + IMPACT 3 (45 %)', 'C18 ปักพื้น · final impact f1 แฟลช · ในเกม hit-stop 100 ms สั่น 7 px flash 140 ms'),
    (115, 'SHAKE (strong)', 'สั่นแรง T115–T122 (4 → 1 px)'),
    (120, 'RING PEAK', 'วงคลื่นกระแทกกว้างสุด'),
    (123, 'RECOVERY_START', 'C19 เริ่มดึงอาวุธขึ้น'),
    (131, 'RESIDUAL', 'final จางลง · residual'),
    (146, 'ANIMATION_END', 'กลับ idle'),
]
LAYERS = [(1, 'ฉากหลัง'), (2, 'ออร่าพื้น + ครึ่งไกลของวง'), (3, 'มังกรส่วนหลังตัว'), (4, 'ตัวละคร + อาวุธ'),
          (5, 'มังกรส่วนหน้า (เว้นหน้าและปลายอาวุธ)'), (6, 'ประกายปลายอาวุธ / จุดสัมผัส / ครึ่งใกล้ของวง')]

TITLE = 'Timeline Sync Sheet — Ultimate (Preview v3)  ·  Bai Long · Nine Dragon Sequence'
SUBTITLE = f'60 FPS · {N} frames (T000–T{N - 1:03d}) = {N / FPS:.2f} s · 1 column = 1 frame (16.7 ms) · ใช้เป็นแม่แบบจังหวะของ Ultimate ตัวละครใหม่'
HITS = {30: 'HIT 1', 45: 'HIT 2', 113: 'HIT 3'}
KEYS = [0, 15, 30, 45, 55, 70, 91, 100, 108, 113, 117, 120, 128, 146]
TRAVEL_LABEL, HEIGHT_LABEL = 'พุ่งไปหน้า (px, 80 = 1 ช่อง)', 'ความสูงลอย (px)'
NOTES = [
    'แถบสีในแถว VFX = ช่วงที่เอฟเฟกต์อยู่บนจอ · ความเข้ม = opacity (จาง = fade-in/out) · ตัวเลข = เฟรมของแผ่น VFX ที่เริ่มแสดงตรงนั้น',
    'กรอบแดงในแถวตัวละคร = เฟรมที่ดาเมจเข้า (Hit 1 T030 · Hit 2 T045 · Hit 3 T113) · C05 ใช้ภาพเดียวกับ C04 (ค้างระหว่างพุ่ง)',
    'ตัวละครวาดอยู่กลางช่องทุกเฟรม — การพุ่ง/การลอยมาจากเส้น "พุ่งไปหน้า" และ "ความสูงลอย" (smoothstep ระหว่างคีย์)',
    'ส่งไม้ต่อ VFX: charge > hint > emerge > ascent > coil > core > dive > final > residual (ซ้อนกัน 2–6 เฟรมเสมอ ไม่มีช่วงว่าง)',
    'final impact วางที่ G = จุดอาวุธแตะพื้น + 1 ช่อง (80 px) = ช่อง main ของพื้นที่ดาเมจ FRONT_CROSS · dive ต้องหายก่อน T112',
]
NAME = 'ult_sync_sheet_v3'
SHAKE_TAB = None
# a timeline JSON with a 'sheet' block (any skill, any hero) replaces the Bai Long tables above
M = D.get('sheet')
if M:
    TITLE, SUBTITLE, NAME = M['title'], M['subtitle'], M.get('name', NAME)
    PHASES = [(a, b, pid, name, tuple(col)) for a, b, pid, name, col, _ in M['phases']]
    PHASE_TH = {pid: th for _, _, pid, _, _, th in M['phases']}
    POSE = {int(k): tuple(v) for k, v in M['poses'].items()}
    GAME_STEP_MS = M['gameSteps']
    TRACKS = [(k, label, tuple(col), src, layers, anchor, note) for k, label, col, src, layers, anchor, note in M['tracks']]
    TRACK_TH = M['trackNames']
    EVENTS = [tuple(e) for e in M['events']]
    LAYERS = [tuple(x) for x in M['layers']]
    HITS = {t: 'HIT' for t in M['hits']}
    KEYS = M['thumbs']
    TRAVEL_LABEL, HEIGHT_LABEL = M['travelLabel'], M['heightLabel']
    NOTES = M['notes']
    SHAKE_TAB = {int(k): tuple(v) for k, v in M['shake'].items()}


def shake(t):
    if SHAKE_TAB is not None:
        return SHAKE_TAB.get(t, (0, 0))
    if 33 <= t <= 36:
        return [(1, 0), (-1, 1), (1, -1), (0, 0)][t - 33]
    if 115 <= t <= 122:
        amp = [4, 4, 3, 3, 2, 2, 1, 1][t - 115]
        return ((-1) ** t * amp, (-1) ** (t // 2) * (amp // 2))
    return (0, 0)


def phase_of(t):
    # a preview moment past the clip (e.g. a buff expiring seconds later) takes the last phase's colour
    return next((p for p in PHASES if p[0] <= t <= p[1]), PHASES[-1])


PF = D['perFrame']
X0, G = PF[0]['root'][0], D['ground']
VFX = {}
for v in D['vfx']:
    VFX.setdefault(v['track'], {})[v['T']] = v
CHAR = [(c['from'], c['to'], int(c['char'][1:])) for c in D['characterTiming']]
travel = [p['root'][0] - X0 for p in PF]
height = [G - p['root'][1] for p in PF]
hits = HITS

# ------------------------------------------------------------------ thumbnails from the MP4
tmp = tempfile.mkdtemp()
subprocess.run(['ffmpeg', '-v', 'error', '-i', MP4, f'{tmp}/f%03d.png'], check=True)
thumbs = {t: Image.open(f'{tmp}/f{t + 1:03d}.png').convert('RGB') for t in KEYS}
if M and M.get('thumbCrop'):          # small characters: zoom the thumbnails on the action
    thumbs = {t: im.crop(tuple(M['thumbCrop'])) for t, im in thumbs.items()}

# ------------------------------------------------------------------ X-sheet PNG
FONT = os.environ.get('SYNC_SHEET_FONT', '/usr/share/fonts/opentype/tlwg/Loma.otf')
FONTB = os.environ.get('SYNC_SHEET_FONT_BOLD', '/usr/share/fonts/opentype/tlwg/Loma-Bold.otf')
f = lambda s, b=False: ImageFont.truetype(FONTB if b else FONT, s)
CW, LX = max(13, 1700 // N), 300                      # px per timeline frame · label column width
W = LX + N * CW + 40
BG, INK, DIM, GRID = (16, 18, 28), (232, 234, 242), (140, 146, 165), (44, 48, 66)
xs = lambda t: LX + t * CW            # left edge of frame t

rows = []                             # (key, height)
TH_W, TH_H = 150, 98
if M and len(M['thumbs']) > 1:            # many preview moments: shrink so they don't overlap
    TH_W = min(TH_W, int((N * max(13, 1700 // N)) / len(M['thumbs'])) - 6)
    TH_H = round(TH_W * 98 / 150)
y = 0
H_TITLE, H_THUMB = 96, TH_H + 70
layout = {}
y = H_TITLE + H_THUMB
LANES = (M or {}).get('lanes') or (9 if len(EVENTS) > 16 else 5)
for key, h in [('ruler', 44), ('phase', 34), ('event', LANES * 28 + 10), ('char', 54), ('travel', 64), ('height', 72)] + \
        [(t[0], 34) for t in TRACKS] + [('shake', 46), ('legend', 250)]:
    layout[key] = (y, h)
    y += h + (6 if key not in ('legend',) else 0)
H = y + 20
img = Image.new('RGB', (W, H), BG)
d = ImageDraw.Draw(img, 'RGBA')

d.text((24, 18), TITLE, font=f(30, True), fill=INK)
d.text((24, 58), SUBTITLE, font=f(18), fill=DIM)

# thumbnails, spread evenly, leader line to their frame
ty = H_TITLE + 6
gap = (N * CW - TH_W) / (len(KEYS) - 1)
ruler_y = layout['ruler'][0]
for i, t in enumerate(KEYS):
    tx = int(LX + i * gap)
    im = thumbs[t].resize((TH_W, TH_H), Image.LANCZOS)
    img.paste(im, (tx, ty))
    ph = phase_of(t)
    d.rectangle([tx - 1, ty - 1, tx + TH_W, ty + TH_H], outline=ph[4], width=2)
    cap = (M or {}).get('thumbCaptions', {}).get(str(t)) or f'T{t:03d}  C{next((c for a, b, c in CHAR if a <= t <= b), CHAR[-1][2]):02d}'
    d.text((tx + 4, ty + TH_H + 4), cap, font=f(15, True), fill=INK)
    if M and M.get('thumbLabels') and str(t) in M['thumbLabels']:   # named preview moments (READY / HIT 1 …)
        tag = M['thumbLabels'][str(t)]
        tw = d.textlength(tag, font=f(13, True))
        d.rectangle([tx + 3, ty + 3, tx + 11 + tw, ty + 22], fill=(10, 10, 16, 200))
        d.text((tx + 7, ty + 5), tag, font=f(13, True), fill=ph[4])
    d.line([(tx + TH_W // 2, ty + TH_H + 26), (xs(min(t, N - 1)) + CW // 2, ruler_y)], fill=ph[4] + (170,), width=2)

# grid + ruler
y0, y1 = layout['ruler'][0], layout['shake'][0] + layout['shake'][1]
for t in range(N + 1):
    col = (70, 76, 100) if t % 10 == 0 else GRID if t % 5 == 0 else (28, 31, 44)
    d.line([(xs(t), y0 + 20), (xs(t), y1)], fill=col, width=1)
STEP_LBL = 10 if CW < 20 else 5
for t in range(0, N, STEP_LBL):
    d.text((xs(t) + 2, y0 + 2), f'T{t:03d}', font=f(13, True), fill=INK)
    d.text((xs(t) + 2, y0 + 22), f'{t / FPS:.2f}s', font=f(11), fill=DIM)
for key, label in [('ruler', 'เวลา (T @60fps)'), ('phase', 'ช่วง (Phase)'), ('event', 'Events / จังหวะโดน'),
                   ('char', 'ท่าตัวละคร C## / เฟรมเกม F##'), ('travel', TRAVEL_LABEL),
                   ('height', HEIGHT_LABEL), ('shake', 'สั่นกล้อง (px)')]:
    yy, hh = layout[key]
    d.text((16, yy + hh // 2 - 11), label, font=f(17, True), fill=INK)

# phases
yy, hh = layout['phase']
for a, b, pid, name, col in PHASES:
    d.rectangle([xs(a), yy, xs(b + 1) - 1, yy + hh], fill=col + (200,))
    d.text((xs(a) + 4, yy + 7), f'{pid}  {PHASE_TH[pid]}', font=f(15, True), fill=(10, 10, 16))

# events: marker + staggered label
yy, hh = layout['event']
lane_end = [-1e9] * LANES
lane_spans = [[] for _ in range(LANES)]
for i, (t, ev, _) in enumerate(EVENTS):
    hit = 'DAMAGE' in ev
    col = (255, 90, 90) if hit else (250, 210, 90) if ev.startswith(('STRIKE', 'FINAL', 'ANTICIPATION')) else (160, 200, 255)
    x = xs(t) + CW // 2
    label = f'T{t:03d} {ev}'
    lw = d.textlength(label, font=f(13, hit))
    lx = x + 4 if x + 4 + lw < W - 10 else x - 4 - lw     # labels near the right edge are drawn to the left
    if M:   # greedy lanes: the first lane where this label's span overlaps nothing already placed
        free = [k for k in range(LANES) if all(lx > b + 6 or lx + lw < a - 6 for a, b in lane_spans[k])]
        lane = free[0] if free else min(range(LANES), key=lambda k: len(lane_spans[k]))
        lane_spans[lane].append((lx, lx + lw))
    else:
        lane = i % LANES
    ly = yy + 6 + lane * 28
    d.line([(x, ly + 10), (x, yy + hh)], fill=col + (200,), width=3 if hit else 1)
    d.polygon([(x - 6, yy + hh - 10), (x + 6, yy + hh - 10), (x, yy + hh)], fill=col)
    d.text((lx, ly), label, font=f(13, hit), fill=col)

# character holds
yy, hh = layout['char']
for a, b, c in CHAR:
    ph = phase_of(a)[4]
    d.rectangle([xs(a) + 1, yy, xs(b + 1) - 2, yy + hh], fill=ph + (95,), outline=ph, width=2)
    d.text((xs(a) + 3, yy + 4), f'C{c:02d}', font=f(13, True), fill=INK)
    d.text((xs(a) + 3, yy + 24), POSE[c][0], font=f(12), fill=DIM)
for t in hits:
    d.rectangle([xs(t), yy - 2, xs(t + 1), yy + hh + 2], outline=(255, 90, 90), width=3)


def curve(key, vals, vmax, col, unit):
    yy, hh = layout[key]
    pts = [(xs(t) + CW / 2, yy + hh - 6 - (v / vmax) * (hh - 14)) for t, v in enumerate(vals)]
    d.line([(xs(0), yy + hh - 6), (xs(N), yy + hh - 6)], fill=GRID, width=1)
    d.line(pts, fill=col, width=3)
    last = None
    for t, v in enumerate(vals):
        r = round(v)
        keyed = t in KEY_TS if M else (r in (14, 18, 22, 30, 58, 64) and last is not None)
        if r != last and (t in (0, N - 1) or keyed):
            d.text((pts[t][0] - 6, pts[t][1] - 18), f'{r}', font=f(12, True), fill=col)
        last = r


KEY_TS = {k for k, _ in D['rootTracks']['travel'] + D['rootTracks']['height']}
curve('travel', travel, max(32, max(travel) + 2), (120, 200, 255), 'px')
curve('height', height, max(66, max(height) + 2), (255, 200, 110), 'px')

# VFX tracks: cell shade = opacity · frame number where it changes · layer tag
for key, label, col, src, layers, anchor, note in TRACKS:
    yy, hh = layout[key]
    d.text((16, yy + 2), f'{label}', font=f(15, True), fill=col)
    d.text((16, yy + 19), layers, font=f(11), fill=DIM)
    tr = VFX.get(key, {})
    last = None
    for t in range(N):
        v = tr.get(t)
        if not v:
            last = None
            continue
        a = int(40 + 215 * v['opacity'])
        d.rectangle([xs(t), yy + 3, xs(t + 1) - 1, yy + hh - 3], fill=col + (a,))
        if v['frame'] != last:
            d.line([(xs(t), yy + 3), (xs(t), yy + hh - 3)], fill=(10, 10, 16), width=2)
            d.text((xs(t) + 2, yy + 8), str(v['frame']), font=f(13, True), fill=(10, 10, 16))
        last = v['frame']

# shake
yy, hh = layout['shake']
mid = yy + hh // 2
d.line([(xs(0), mid), (xs(N), mid)], fill=GRID)
for t in range(N):
    sx, sy = shake(t)
    amp = max(abs(sx), abs(sy))
    if amp:
        k = min(5, (hh / 2 - 3) / max(max(max(abs(a), abs(b)) for a, b in (shake(x) for x in range(N))), 1))
        d.rectangle([xs(t) + 2, mid - amp * k, xs(t + 1) - 2, mid + amp * k], fill=(255, 110, 110, 220))
        peak = SHAKE_TAB is not None and amp == max(max(abs(a), abs(b)) for a, b in SHAKE_TAB.values())
        if t in (33, 115, 119, 121) and SHAKE_TAB is None or peak:
            d.text((xs(t + 1) + 3, yy + 1), f'{amp}px', font=f(12, True), fill=(255, 140, 140))

# legend
yy, hh = layout['legend']
d.text((24, yy + 6), 'อ่านยังไง' if not M else 'Implementation notes', font=f(20, True), fill=INK)
lines = NOTES
for i, s in enumerate(lines):
    d.text((24, yy + 42 + i * 26), '• ' + s, font=f(15), fill=INK)
lx = 24
ly = yy + 42 + len(lines) * 26 + 14
d.text((lx, ly), 'ชั้นภาพ (หลัง > หน้า):', font=f(15, True), fill=INK)
lx += 190
for n, name in LAYERS:
    s = f'L{n} {name}'
    d.text((lx, ly), s, font=f(14), fill=DIM)
    lx += d.textlength(s, font=f(14)) + 26

os.makedirs(OUT, exist_ok=True)
img.save(f'{OUT}/{NAME}.png', optimize=True)

# ------------------------------------------------------------------ workbook
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
AR = 'Arial'
HEAD = PatternFill('solid', fgColor='1F2433')
hfont = Font(name=AR, bold=True, color='FFFFFF')
thin = Side(style='thin', color='D0D4DE')
hex_ = lambda c: '%02X%02X%02X' % c
light = lambda c, k=.6: hex_(tuple(int(v + (255 - v) * k) for v in c))


def header(ws, cols, widths):
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HEAD, hfont
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = 'B2'


# Read Me
ws = wb.active
ws.title = 'Read Me'
ws.column_dimensions['A'].width = 120
README = M['readme'] if M and 'readme' in M else [
    'Timeline Sync Sheet — Ultimate (Preview v3) · Bai Long · Nine Dragon Sequence',
    '',
    'ใช้เป็นแม่แบบจังหวะของ Ultimate ตัวละครใหม่: 60 FPS, 147 เฟรม (T000–T146) = 2.45 วินาที',
    'Frames = ทุกเฟรม: ช่วง, ท่าตัวละคร, ตำแหน่ง, VFX แต่ละ track (เฟรม/opacity/การแบ่งชั้น), event, สั่นกล้อง',
    'Character = 19 ท่า (C01–C20, C05 ใช้ภาพ C04) + ช่วงเวลาใน preview + ความยาวในเกม',
    'VFX Tracks = แผ่นที่ใช้ จุดยึด ชั้น และหมายเหตุของแต่ละเอฟเฟกต์ · Events = จังหวะสำคัญ · Layers = ลำดับชั้นภาพ',
    '',
    'หน่วย: px = พิกเซลของแผ่นเกม (canvas 460×300, พื้น y = 250) · 80 px = 1 ช่องสนาม ≈ ความสูงตัวละครตอนยืน',
    'opacity: 1 = เต็ม · การแบ่งชั้น (split): lower = ครึ่งไกลหลังตัว/ครึ่งใกล้หน้าตัว · head40 = หัว 40 % หน้าตัว ลำตัว 60 % หลังตัว · '
    'tip = ปลายแสงอยู่หน้า · band = ลำตัวพาดผ่าน · head-bottom = หัวด้านล่างอยู่หน้า · front = หน้าสุด',
    'ที่มา: bailong_ult_v3_timeline.json + tools/sprites/bailong_ult_preview.py (preview v3) · ไฟล์นี้สร้างจาก tools/sprites/ult_sync_sheet.py',
]
for r, s in enumerate(README, 1):
    c = ws.cell(row=r, column=1, value=s)
    c.font = Font(name=AR, bold=(r == 1), size=14 if r == 1 else 11)
    c.alignment = Alignment(wrap_text=True, vertical='top')

# Frames
ws = wb.create_sheet('Frames')
tcols = [t[0] for t in TRACKS]
cols = ['T', 'ms', 'Phase', 'Char', 'Game frame', 'Hold #', 'Travel px', 'Height px', 'Travel cells', 'Weapon tip x', 'Weapon tip y',
        'Body x', 'Body y', 'Event'] + [f'{t}\n{TRACK_TH[t]}' for t in tcols] + ['Shake dx', 'Shake dy', 'Note']
header(ws, cols, [6, 8, 18, 7, 8, 7, 9, 9, 9, 9, 9, 8, 8, 30] + [13] * len(tcols) + [8, 8, 34])
ev_at = {}
for t, ev, note in EVENTS:
    ev_at.setdefault(t, []).append((ev, note))
for t in range(N):
    p = PF[t]
    c = int(p['char'][1:])
    a = next(a for a, b, cc in CHAR if cc == c)
    ph = phase_of(t)
    r = t + 2
    vf = []
    for k in tcols:
        v = VFX.get(k, {}).get(t)
        vf.append(f"f{v['frame']} · {round(v['opacity'] * 100)}% · {v['split']}" if v else '')
    sx, sy = shake(t)
    evs = ev_at.get(t, [])
    ws.append([t, f'=A{r}*1000/60', f'{ph[2]} {PHASE_TH[ph[2]]}', p['char'], POSE[c][0], t - a + 1,
               round(travel[t], 1), round(height[t], 1), f'=G{r}/80', p['spearTip'][0], p['spearTip'][1],
               p['center'][0], p['center'][1], ' / '.join(e for e, _ in evs)] + vf + [sx, sy, ' / '.join(n for _, n in evs)])
    fill = PatternFill('solid', fgColor=light(ph[4], .72))
    for ci in range(1, len(cols) + 1):
        cell = ws.cell(row=r, column=ci)
        cell.font = Font(name=AR, size=10, bold=bool(evs and 'DAMAGE' in evs[0][0]))
        cell.border = Border(bottom=thin)
        if ci <= 6:
            cell.fill = fill
    for j, k in enumerate(tcols):
        if VFX.get(k, {}).get(t):
            col = next(tt[2] for tt in TRACKS if tt[0] == k)
            ws.cell(row=r, column=15 + j).fill = PatternFill('solid', fgColor=light(col, .45))
    ws.cell(row=r, column=2).number_format = '0.0'
    ws.cell(row=r, column=9).number_format = '0.00'
    if t in hits:
        for ci in range(1, len(cols) + 1):
            ws.cell(row=r, column=ci).border = Border(top=Side(style='medium', color='D04040'), bottom=Side(style='medium', color='D04040'))

# Character
ws = wb.create_sheet('Character')
header(ws, ['Char', 'Game frame', 'From T', 'To T', 'Frames', 'Preview ms', 'Game ms', 'ท่า', 'จุดสำคัญ'],
       [7, 10, 8, 8, 8, 11, 9, 46, 46])
gi = 0
for i, (a, b, c) in enumerate(CHAR):
    r = i + 2
    game = ''
    if (c != 5 or M) and gi < len(GAME_STEP_MS):
        game = GAME_STEP_MS[gi]
        gi += 1
    ws.append([f'C{c:02d}', POSE[c][0], a, b, f'=D{r}-C{r}+1', f'=E{r}*1000/60', game, POSE[c][1], POSE[c][2]])
    ws.cell(row=r, column=6).number_format = '0'
    for ci in range(1, 10):
        ws.cell(row=r, column=ci).font = Font(name=AR, size=10)
        ws.cell(row=r, column=ci).fill = PatternFill('solid', fgColor=light(phase_of(a)[4], .72))
n = len(CHAR) + 2
ws.cell(row=n, column=4, value='รวม').font = Font(name=AR, bold=True)
ws.cell(row=n, column=5, value=f'=SUM(E2:E{n - 1})').font = Font(name=AR, bold=True)
ws.cell(row=n, column=6, value=f'=SUM(F2:F{n - 1})').font = Font(name=AR, bold=True)
ws.cell(row=n, column=7, value=f'=SUM(G2:G{n - 1})').font = Font(name=AR, bold=True)
ws.cell(row=n, column=6).number_format = '0'
ws.cell(row=n + 1, column=1, value=M.get('gameNote', 'Game ms = ความยาวของเฟรมนั้นในเกม') if M else 'Game ms = ความยาวของเฟรมนั้นในเกม (SHEETS hero-bailong.png ult.steps) — C05 ไม่มีในเกม (รวมอยู่ใน F4)').font = Font(name=AR, italic=True, size=9)

# VFX tracks
ws = wb.create_sheet('VFX Tracks')
header(ws, ['Track', 'ชื่อ', 'แผ่น (เฟรม)', 'แสดง T', 'เฟรมที่ใช้', 'ชั้น', 'จุดยึด', 'หมายเหตุ'], [10, 20, 30, 12, 20, 34, 40, 60])
for i, (key, label, col, src, layers, anchor, note) in enumerate(TRACKS):
    tr = VFX.get(key, {})
    ts = sorted(tr)
    runs, start = [], None
    for t in ts:
        if start is None:
            start = prev = t
        elif t != prev + 1:
            runs.append((start, prev))
            start = t
        prev = t
    if start is not None:
        runs.append((start, prev))
    frames = sorted({v['frame'] for v in tr.values()})
    ws.append([key, TRACK_TH[key], src, ', '.join(f'{a:03d}–{b:03d}' for a, b in runs), ', '.join(map(str, frames)), layers, anchor, note])
    for ci in range(1, 9):
        c = ws.cell(row=i + 2, column=ci)
        c.font = Font(name=AR, size=10)
        c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(row=i + 2, column=1).fill = PatternFill('solid', fgColor=light(col, .35))

# Events
ws = wb.create_sheet('Events')
header(ws, ['T', 'ms', 'Event', 'หมายเหตุ'], [7, 9, 28, 90])
for i, (t, ev, note) in enumerate(EVENTS):
    r = i + 2
    ws.append([t, f'=A{r}*1000/60', ev, note])
    for ci in range(1, 5):
        ws.cell(row=r, column=ci).font = Font(name=AR, size=10, bold='DAMAGE' in ev, color='C03030' if 'DAMAGE' in ev else '000000')
    ws.cell(row=r, column=2).number_format = '0'

# Layers
ws = wb.create_sheet('Layers')
header(ws, ['Layer', 'เนื้อหา'], [8, 70])
for n_, name in LAYERS:
    ws.append([f'L{n_}', name])
    for ci in (1, 2):
        ws.cell(row=ws.max_row, column=ci).font = Font(name=AR, size=11)
ws.append([])
layer_note = 'หน้าตัวละครกับปลายอาวุธถูกเว้นออกจาก L5 ทุกเฟรม (เอฟเฟกต์ไม่บังหน้า/ปลายอาวุธ)'
ws.append(['', M.get('layerNote', layer_note) if M else layer_note])
ws.cell(row=ws.max_row, column=2).font = Font(name=AR, italic=True)

wb.save(f'{OUT}/{NAME}.xlsx')
print('ok', img.size)
